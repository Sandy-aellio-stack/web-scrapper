import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.robotparser import RobotFileParser

import aqua_scraper as s
from docx import Document

HTML = '''<!doctype html><html><head><title>Requirements guide</title>
<meta name="description" content="A guide to requirements"></head><body>
<!-- comments must not crash extraction -->
<header><p>Header text which must not enter the report</p></header>
<div class="mob-menu-container"><p>Mobile navigation must be removed as well</p></div>
<main><h1>Requirements guide</h1><p>Write a clear requirement with an expected behaviour before creating tests.</p>
<h2>Coverage</h2><p>Connect each requirement to relevant test cases and review missing coverage.</p>
<p>Connect each requirement to relevant test cases and review missing coverage.</p>
<div>A long leaf div is a real content block that can describe how requirements connect to tests and defects in a project.</div>
<ul><li><p>Review every linked test case before approving this requirement for execution.</p></li></ul>
<table><tr><th>Role</th><th>Permission</th></tr><tr><td>Reviewer</td><td>Read requirements and record review findings</td></tr></table>
<a href="/pricing/?utm_source=test#plans">Pricing</a><a href="https://evil.example/">External</a>
<iframe src="https://tour.example/view"></iframe></main>
<footer><p>Footer text which must not enter the report</p></footer></body></html>'''


class ScraperTests(unittest.TestCase):
    def test_url_scope(self):
        self.assertIsNone(s.normalise('https://aqua-cloud.io.evil.example/pricing/'))
        self.assertIsNone(s.normalise('https://app.aqua-cloud.io/'))
        self.assertIsNone(s.normalise('https://aqua-cloud.io:8443/pricing/'))
        self.assertIsNone(s.normalise('javascript:alert(1)'))
        self.assertIsNone(s.normalise('https://me:password@aqua-cloud.io/'))

    def test_normalisation(self):
        self.assertEqual(s.normalise('/pricing/?utm_source=abc#plans'), 'https://aqua-cloud.io/pricing/')

    def test_scope_exclusions(self):
        for url in ['/tag/tests/', '/wp-admin/', '/guide.pdf', '/de/pricing/', '/aqua-vs-testrail/']:
            self.assertIsNone(s.normalise(url))

    def test_document_blocks_and_chrome(self):
        data = s.extract_html(HTML, s.SEEDS[0])
        text = '\n'.join(b['text'] for b in data['blocks'])
        self.assertNotIn('Header text', text)
        self.assertNotIn('Footer text', text)
        self.assertNotIn('Mobile navigation', text)
        self.assertIn('A long leaf div', text)
        self.assertEqual(text.count('Connect each requirement'), 1)
        self.assertEqual(text.count('Review every linked'), 1)
        self.assertIn('Reviewer | Read requirements', text)

    def test_links_and_embeds(self):
        data = s.extract_html(HTML, s.SEEDS[0])
        self.assertIn('https://aqua-cloud.io/pricing/', data['links'])
        self.assertFalse(any('evil.example' in u for u in data['links']))
        self.assertEqual(data['embeds'], ['https://tour.example/view'])

    def test_availability_labels_are_not_lost(self):
        data = s.extract_html('<html><body><span>COMING SOON</span><h2>Project AI</h2><p>This feature is described as planned rather than available.</p></body></html>', s.SEEDS[0])
        self.assertEqual(data['blocks'][0]['text'], 'COMING SOON')

    def test_topic_routing(self):
        self.assertEqual(s.TOPICS[s.topic_for('https://aqua-cloud.io/pricing/')][0], 'Pricing licences and support')
        self.assertEqual(s.TOPICS[s.topic_for('https://aqua-cloud.io/requirements-engineering-tool/')][0], 'Requirements and traceability')

    def test_gitbook_markdown(self):
        raw = '# Getting started\n\nThis page explains how projects organise the test cases and requirements.\n\n[General](/documentation/web-client/general.md)\n\n# Agent Instructions\n\nUnrelated platform instructions.'
        data, method = s.extract_response(raw, 'https://docs.aqua-cloud.io/documentation', 'text/plain; charset=utf-8')
        self.assertEqual(data['title'], 'Getting started')
        self.assertEqual(method, 'http Markdown')
        self.assertIn('https://docs.aqua-cloud.io/documentation/web-client/general.md', data['links'])
        self.assertNotIn('Unrelated platform', str(data['blocks']))

    def test_browser_failure_marks_thin_page_as_issue(self):
        with tempfile.TemporaryDirectory() as tmp:
            args = s.arguments(['--max-pages', '1', '--output', tmp])
            with patch.object(s.Fetcher, 'allowed', return_value=True), patch.object(s.Fetcher, 'raw', return_value=('<h1>Loading</h1>', s.SEEDS[0], 200, 'text/html')), patch.object(s.BrowserRenderer, 'render', side_effect=s.FetchError('Browser unavailable')):
                self.assertEqual(s.crawl(args), 2)
            bundle = json.loads((Path(tmp) / 'aqua_data.json').read_text())
            self.assertIn('Browser unavailable', bundle['issues'][0]['reason'])

    def test_short_markdown_is_not_mistaken_for_spa(self):
        raw = '# General\n\nProjects contain requirements, test cases and defects. Use the linked guides to understand each item and the available operations.\n\n[Details](/documentation/web-client/test-cases.md)'
        url = 'https://docs.aqua-cloud.io/documentation/web-client/general.md'
        with tempfile.TemporaryDirectory() as tmp:
            args = s.arguments(['--seed', url, '--max-pages', '1', '--output', tmp])
            with patch.object(s.Fetcher, 'allowed', return_value=True), patch.object(s.Fetcher, 'raw', return_value=(raw, url, 200, 'text/plain')), patch.object(s.BrowserRenderer, 'render', side_effect=AssertionError('Markdown did not need a browser')):
                self.assertEqual(s.crawl(args), 0)
            bundle = json.loads((Path(tmp) / 'aqua_data.json').read_text())
            self.assertIn('https://docs.aqua-cloud.io/documentation/web-client/test-cases.md', bundle['pending_urls'])

    def test_robots_deny(self):
        f = s.Fetcher(delay=0)
        with patch.object(f, 'raw', return_value=('User-agent: *\nDisallow: /private/', '', 200, 'text/plain')):
            self.assertFalse(f.allowed('https://aqua-cloud.io/private/page'))
            self.assertTrue(f.allowed(s.SEEDS[0]))

    def test_robots_unavailable_fails_closed(self):
        f = s.Fetcher(delay=0)
        with patch.object(f, 'raw', side_effect=s.FetchError('HTTP 503')):
            self.assertFalse(f.allowed(s.SEEDS[0]))
            self.assertIn('unavailable', f.robot_status['aqua-cloud.io'])

    def test_robots_404(self):
        f = s.Fetcher(delay=0)
        with patch.object(f, 'raw', side_effect=s.FetchError('HTTP 404')):
            self.assertTrue(f.allowed(s.SEEDS[0]))

    def test_external_redirect_rejected(self):
        f = s.Fetcher(delay=0, retries=0)
        e = HTTPError(s.SEEDS[0], 302, 'redirect', {'Location': 'https://evil.example/private'}, None)
        with patch.object(f.opener, 'open', side_effect=e):
            with self.assertRaisesRegex(s.FetchError, 'outside allowed'):
                f.raw(s.SEEDS[0])

    def test_crawl_word_and_offline_rebuild(self):
        with tempfile.TemporaryDirectory() as tmp:
            args = s.arguments(['--mode', 'http', '--max-pages', '1', '--output', tmp])
            with patch.object(s.Fetcher, 'allowed', return_value=True), patch.object(s.Fetcher, 'raw', return_value=(HTML, s.SEEDS[0], 200, 'text/html')):
                self.assertEqual(s.crawl(args), 0)
            bundle = json.loads((Path(tmp) / 'aqua_data.json').read_text())
            self.assertEqual(len(bundle['pages']), 1)
            self.assertGreater(bundle['pending_count'], 0)
            doc = Document(Path(tmp) / 'Aqua_Cloud_Handbook.docx')
            self.assertIn('Aqua Cloud Product Handbook', doc.paragraphs[0].text)
            with patch.object(s.Fetcher, 'raw', side_effect=AssertionError('offline rebuild used network')):
                self.assertEqual(s.main(['--from-json', str(Path(tmp) / 'aqua_data.json'), '--output', tmp]), 0)

    def test_empty_crawl_does_not_create_handbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            args = s.arguments(['--mode', 'http', '--max-pages', '1', '--output', tmp])
            with patch.object(s.Fetcher, 'allowed', return_value=True), patch.object(s.Fetcher, 'raw', return_value=('<html><body><h1>Loading</h1></body></html>', s.SEEDS[0], 200, 'text/html')):
                self.assertEqual(s.crawl(args), 2)
            self.assertFalse((Path(tmp) / 'Aqua_Cloud_Handbook.docx').exists())
            self.assertTrue((Path(tmp) / 'aqua_data.json').exists())

    def test_word_limit_label(self):
        data = s.extract_html(HTML, s.SEEDS[0])
        data.update(id='S001', topic=0, method='fixture', retrieved_at='test', notes=[])
        bundle = dict(pages=[data], issues=[], collected_at='test', attempted=1, pending_count=0)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'test.docx'; s.write_report(bundle, path, max_words=10)
            text = '\n'.join(p.text for p in Document(path).paragraphs)
            self.assertIn('truncated by the configured word limit', text)


if __name__ == '__main__':
    unittest.main()
