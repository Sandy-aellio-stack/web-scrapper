#!/usr/bin/env python3
"""Collect public aqua documentation and produce a source-organised Word handbook.

Python 3.10+. No account, API key, or language model required.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
import sys
import time
from collections import deque
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlsplit, urlunsplit
from urllib.request import Request, build_opener, HTTPRedirectHandler
from urllib.robotparser import RobotFileParser

from lxml import html, etree
from markdown_it import MarkdownIt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

VERSION = "1.0.0"
USER_AGENT = "AquaDocumentationCollector/1.0"
HOSTS = {"aqua-cloud.io", "www.aqua-cloud.io", "docs.aqua-cloud.io"}
SEEDS = ["https://aqua-cloud.io/product-tour/",
         "https://docs.aqua-cloud.io/documentation",
         "https://aqua-cloud.io/aqua-cloud-faq/",
         "https://aqua-cloud.io/pricing/"]
SKIP_PATH = re.compile(r"/(?:de|es|fr|blog|tag|category|author|wp-admin|wp-json|wp-content|contact-page|privacy-policy|imprint|terms-and-conditions)(?:/|$)|/aqua-vs-", re.I)
FILE_SUFFIX = re.compile(r"\.(?:pdf|png|jpe?g|webp|gif|svg|zip|mp4|mp3|css|js|xml|ico|woff2?)$", re.I)
TOPICS = [
    ("Product overview", ("product-tour", "getting-started", "aqua-cloud-faq", "about-us"),
     "Start here to understand the product purpose and terminology. The collected pages below describe what the vendor offers. Performance and savings statements are vendor claims, not independently measured results."),
    ("Requirements and traceability", ("requirement", "traceab", "dependenc", "coverage"),
     "Requirements describe the behaviour a system needs to deliver. Traceability connects those requirements with tests and defects so a team can investigate coverage and the impact of a change. Read the source details for the specific relationships supported by aqua."),
    ("Test cases and test data", ("test-case", "test-data", "parameter", "test-values", "test-management"),
     "A test case describes an action and its expected result. Test data supplies the values used when running it. These pages explain available structures, editing operations and reuse options; they should not be treated as a guarantee that every option is included in every subscription."),
    ("Planning execution and results", ("execution", "test-run", "test-scenar", "test-plan", "test-cycle", "test-step"),
     "Planning decides which tests to run, who will run them and what environment is required. Execution records actual outcomes. Read the source instructions to understand how these records are connected and how execution results can be reviewed."),
    ("Defects and Capture", ("bug", "defect", "capture"),
     "A defect record explains a problem, the conditions that reproduce it and the difference between expected and actual behaviour. Capture-related pages describe recording capabilities. External Capture documentation is outside the default crawl scope."),
    ("Automation and continuous delivery", ("automation", "automated", "jenkins", "selenium", "ranorex", "jmeter", "uft", "soapui"),
     "Automation integrations connect external execution tools with test management. Distinguish a connector that imports results from a tool that executes tests itself. Check the collected setup instructions before assuming an integration provides both."),
    ("AI capabilities", ("intelligence", "copilot", "artificial", "/ai", "aqua-ai"),
     "AI-assisted generation can help prepare testing material, but generated content still needs review against the requirements. The source material below may include released features, previews and planned features. Preserve their availability labels when evaluating them."),
    ("Integrations and API", ("integrat", "rest-api", "jira", "azure", "confluence", "powershell", "unixshell", "database"),
     "Integration documentation explains how data moves between systems. Look for supported object types, synchronisation direction, authentication, setup prerequisites and limits. A marketing integration listing alone does not establish those technical details."),
    ("Agile work and collaboration", ("agile", "scrum", "sprint", "backlog", "workflow", "discussion", "notification"),
     "These pages relate testing to team planning and day-to-day collaboration. Workflows describe state transitions; boards and sprints organise work. Check how tests, requirements and defects participate in each workflow."),
    ("Reports dashboards and history", ("report", "dashboard", "histor", "audit", "logging"),
     "Reports and dashboards summarise recorded information. History and audit records help explain how information changed. Useful evaluation questions include which fields are available, how filters work and whether the result can be exported."),
    ("Administration security and deployment", ("secur", "permission", "role", "user", "authentication", "sso", "data-center", "hosting", "deploy", "install", "administr", "project", "general"),
     "Administration covers users, projects, access and deployment. Security claims require careful reading of scope and conditions. A statement about data protection should not be interpreted as evidence of a specific certification unless the source explicitly identifies it."),
    ("Migration import and export", ("migrat", "import", "export"),
     "Migration transfers existing records into another system. Verify which objects, attachments, links and history are retained. Vendor migration-time examples are not a delivery estimate for a particular dataset."),
    ("Pricing licences and support", ("pricing", "licen", "support", "subscription"),
     "Pricing and licence terms can change. Treat the captured text as a dated reference and confirm currency, billing period, included features and contract conditions with the vendor before making a purchase decision."),
    ("Other product documentation", (),
     "Additional relevant pages are collected here when a URL does not map clearly to an earlier topic. This section may include information that also applies to several other chapters."),
]


def clean(text):
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text or "")
    return re.sub(r"\s+", " ", text).strip()


def normalise(url, base=None):
    """Keep public HTML URLs on the explicit host allowlist; drop tracking queries."""
    p = urlsplit(urljoin(base or SEEDS[0], url))
    if p.scheme not in {"http", "https"} or p.hostname not in HOSTS or p.username or p.password or p.port not in {None, 80, 443}:
        return None
    path = re.sub(r"/{2,}", "/", p.path or "/")
    if FILE_SUFFIX.search(path) or SKIP_PATH.search(path):
        return None
    if p.hostname == "docs.aqua-cloud.io" and not path.startswith("/documentation"):
        return None
    return urlunsplit(("https", p.hostname, path, "", ""))


def topic_for(url):
    path = urlsplit(url).path.lower()
    for i, (_, keys, _) in enumerate(TOPICS):
        if any(k in path for k in keys):
            return i
    return len(TOPICS) - 1


def relevant(url):
    return urlsplit(url).hostname == "docs.aqua-cloud.io" or topic_for(url) < len(TOPICS) - 1


def extract_html(raw, url):
    """Extract document-order blocks, not a list of guessed product-tour CSS classes."""
    root = html.fromstring(raw)
    title = clean(" ".join(root.xpath("//title/text()"))) or url
    canonical = root.xpath("//link[@rel='canonical']/@href")
    links = sorted({u for href in root.xpath("//a[@href]/@href") if (u := normalise(href, url))})
    embeds = [urljoin(url, src) for src in root.xpath("//iframe/@src")]
    description = root.xpath("//meta[@name='description']/@content")
    for el in root.xpath("//script|//style|//noscript|//nav|//header|//footer|//form|//svg|//template|//aside|//*[@role='navigation']"):
        if el.getparent() is not None:
            el.drop_tree()
    # Remove identifiable page chrome without hiding content or changing consent.
    for el in list(root.iter()):
        if not isinstance(el.tag, str):
            continue
        marker = (el.get("id", "") + " " + el.get("class", "")).lower()
        if re.search(r"(?:^|[\s_-])(?:cookie-banner|cookiebot|cookieyes|onetrust-banner-sdk|aqua-consent|site-menu|mob-menu-container|header-menu-mobile|mobile-menu)(?:$|[\s_-])", marker):
            if el.getparent() is not None:
                el.drop_tree()
    candidates = root.xpath("//main|//*[@role='main']|//article")
    body = max(candidates, key=lambda e: len(e.text_content())) if candidates else root
    blocks, seen = [], set()
    for el in body.iter():
        tag = el.tag if isinstance(el.tag, str) else ""
        # Preserve short availability labels that would otherwise be lost by
        # the prose length filter, especially unreleased or preview features.
        if tag in {"span", "div", "p", "strong"}:
            label = clean(el.text_content())
            if len(label) <= 35 and re.search(r"\b(?:coming soon|preview|beta|deprecated)\b", label, re.I):
                if label not in seen:
                    blocks.append({"kind": "text", "level": 0, "text": label})
                    seen.add(label)
                continue
        if tag not in {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "dt", "dd", "summary", "tr", "pre", "div"}:
            continue
        if tag == "div" and (len(el) == 0 or not any(isinstance(c.tag, str) and c.tag in {"div", "p", "li", "h1", "h2", "h3", "h4", "table", "ul", "ol"} for c in el)):
            # Leaf divs matter on this site; require substantial prose.
            text = clean(el.text_content())
            if len(text) < 80 or len(el.xpath(".//p|.//li|.//h1|.//h2|.//h3|.//h4|.//table|.//div")):
                continue
        elif tag == "div":
            continue
        if tag in {"p", "li"} and any(a.tag in {"li", "tr", "pre"} for a in el.iterancestors()):
            continue
        if tag == "tr":
            text = " | ".join(clean(x.text_content()) for x in el.xpath("./th|./td"))
        else:
            text = clean(el.text_content())
        if not text or text in seen:
            continue
        if text.lower() in {"learn more", "request a demo", "contact sales", "start free trial", "try it yourself!", "see all integrations"}:
            continue
        if tag not in {"h1", "h2", "h3", "h4", "h5", "h6"} and len(text) < 25:
            continue
        seen.add(text)
        blocks.append({"kind": "heading" if re.fullmatch(r"h[1-6]", tag) else "text", "level": int(tag[1]) if re.fullmatch(r"h[1-6]", tag) else 0, "text": text})
    text = "\n".join(b["text"] for b in blocks)
    return {"title": title, "url": url, "canonical": canonical[0] if canonical else None,
            "description": description[0] if description else "", "blocks": blocks,
            "links": links, "embeds": embeds, "characters": len(text),
            "content_hash": hashlib.sha256(text.encode()).hexdigest()}


def extract_response(raw, url, content_type):
    if any(t in content_type.lower() for t in ["text/html", "application/xhtml+xml"]):
        return extract_html(raw, url), "http HTML"
    if urlsplit(url).hostname == "docs.aqua-cloud.io" and any(t in content_type.lower() for t in ["text/plain", "text/markdown"]):
        # GitBook publishes Markdown to programmatic clients. Its platform-specific
        # agent appendix is not product documentation and is never executed.
        raw = re.split(r"(?m)^# Agent Instructions\s*$", raw, maxsplit=1)[0]
        rendered = MarkdownIt("commonmark").enable("table").render(raw)
        parsed = extract_html("<html><body><main>" + rendered + "</main></body></html>", url)
        headings = [b['text'] for b in parsed['blocks'] if b['kind'] == 'heading']
        parsed['title'] = headings[0] if headings else url
        return parsed, "http Markdown"
    raise FetchError(f"Unsupported content type skipped: {content_type}")


class FetchError(Exception):
    pass


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):
        return None


class Fetcher:
    def __init__(self, delay=1.0, timeout=30, retries=2):
        self.delay, self.timeout, self.retries = delay, timeout, retries
        self.opener = build_opener(NoRedirect())
        self.robots, self.last_request, self.robot_status = {}, {}, {}

    def pace(self, url):
        host = urlsplit(url).hostname
        parser = self.robots.get(host)
        crawl_delay = parser.crawl_delay(USER_AGENT) if parser else None
        delay = max(self.delay, crawl_delay or 0)
        time.sleep(max(0, delay - (time.monotonic() - self.last_request.get(host, 0))))
        self.last_request[host] = time.monotonic()

    def raw(self, url, check_redirects=True):
        current = url
        for _ in range(6):
            for attempt in range(self.retries + 1):
                self.pace(current)
                try:
                    with self.opener.open(Request(current, headers={"User-Agent": USER_AGENT, "Accept": "text/html,text/plain;q=0.9,*/*;q=0.1"}), timeout=self.timeout) as response:
                        content_type = response.headers.get("Content-Type", "")
                        data = response.read(8_000_001)
                        if len(data) > 8_000_000:
                            raise FetchError("Response exceeds 8 MB limit")
                        charset = response.headers.get_content_charset() or "utf-8"
                        return data.decode(charset, errors="replace"), current, response.status, content_type
                except HTTPError as e:
                    if e.code in {301, 302, 303, 307, 308}:
                        target = urljoin(current, e.headers.get("Location", ""))
                        valid = normalise(target) if check_redirects else target
                        if not valid or urlsplit(valid).hostname not in HOSTS:
                            raise FetchError("Redirect outside allowed public scope") from e
                        if check_redirects and not self.allowed(valid):
                            raise FetchError("Redirect blocked by robots policy") from e
                        current = valid
                        break
                    if e.code == 429 or 500 <= e.code < 600:
                        if attempt < self.retries:
                            retry_after = e.headers.get("Retry-After", "")
                            delay = int(retry_after) if retry_after.isdigit() else 2 ** (attempt + 1)
                            time.sleep(min(60, max(1, delay)))
                            continue
                    raise FetchError(f"HTTP {e.code}") from e
                except (URLError, TimeoutError, OSError) as e:
                    if attempt < self.retries:
                        time.sleep(2 ** (attempt + 1))
                        continue
                    raise FetchError(f"{type(e).__name__}: {e}") from e
            else:
                raise FetchError("Request retries exhausted")
        raise FetchError("Too many redirects")

    def allowed(self, url):
        host = urlsplit(url).hostname
        if host not in self.robots:
            try:
                body, _, _, _ = self.raw(f"https://{host}/robots.txt", check_redirects=False)
                if "<html" in body[:500].lower():
                    raise FetchError("robots.txt returned HTML")
                p = RobotFileParser()
                p.parse(body.splitlines())
                self.robots[host] = p
                self.robot_status[host] = "loaded"
            except FetchError as e:
                if str(e) == "HTTP 404":
                    p = RobotFileParser(); p.parse(["User-agent: *", "Allow: /"])
                    self.robots[host] = p
                    self.robot_status[host] = "not published HTTP 404"
                else:
                    self.robots[host] = None
                    self.robot_status[host] = f"unavailable; host skipped: {e}"
        parser = self.robots[host]
        return parser is not None and parser.can_fetch(USER_AGENT, url)


class BrowserRenderer:
    def __init__(self, args, fetcher):
        self.args, self.fetcher, self.manager, self.browser, self.context = args, fetcher, None, None, None

    def start(self):
        if self.context:
            return
        try:
            from playwright.sync_api import sync_playwright
            self.manager = sync_playwright().start()
            self.browser = self.manager.chromium.launch(headless=not self.args.headed)
            self.context = self.browser.new_context(viewport={"width": 1440, "height": 1000}, user_agent=USER_AGENT)
        except Exception as e:
            self.close()
            raise FetchError("Browser unavailable. Run python -m playwright install chromium. " + str(e).splitlines()[0]) from e

    def render(self, url):
        self.start()
        page = self.context.new_page()
        network, notes, snapshots = [], [], []

        def route_handler(route):
            request = route.request
            if request.is_navigation_request():
                target = normalise(request.url)
                if not target or not self.fetcher.allowed(target):
                    route.abort()
                    return
            route.continue_()

        def response_handler(response):
            if response.request.resource_type in {"xhr", "fetch"} and len(network) < 150:
                # Metadata only: do not save tokens, request bodies, cookies or API response bodies.
                p = urlsplit(response.url)
                network.append({"url": urlunsplit((p.scheme, p.netloc, p.path, "", "")), "status": response.status,
                                "content_type": response.headers.get("content-type", "")})

        page.route("**/*", route_handler)
        page.on("response", response_handler)
        try:
            response = page.goto(url, wait_until="domcontentloaded", timeout=self.args.timeout * 1000)
            if response and response.status >= 400:
                raise FetchError(f"Browser HTTP {response.status}")
            try:
                page.wait_for_function("document.body && document.body.innerText.trim().length > 400", timeout=12000)
            except Exception:
                notes.append("Rendered body did not reach content threshold within 12 seconds")
            # Only an explicitly configured consent choice is clicked. Never hide a banner.
            if self.args.consent_selector and page.locator(self.args.consent_selector).count():
                page.locator(self.args.consent_selector).first.click(timeout=5000)
                page.wait_for_timeout(700)
            for _ in range(4):
                page.evaluate("window.scrollBy(0, window.innerHeight)")
                page.wait_for_timeout(300)
            snapshots.append(page.content())
            # Optional inspected selector. It is never guessed from arbitrary Next links.
            if self.args.next_selector:
                for _ in range(self.args.max_steps - 1):
                    locator = page.locator(self.args.next_selector).first
                    if not locator.count() or not locator.is_visible() or not locator.is_enabled():
                        break
                    previous = page.locator("body").inner_text()
                    locator.click(timeout=5000)
                    try:
                        page.wait_for_function("old => document.body.innerText !== old", arg=previous, timeout=5000)
                    except Exception:
                        notes.append("Next step did not change body text; traversal stopped")
                        break
                    page.wait_for_timeout(300)
                    current = page.content()
                    if current in snapshots:
                        notes.append("Repeated tour state; traversal stopped")
                        break
                    snapshots.append(current)
                else:
                    notes.append("Tour step limit reached; later states may be missing")
            return snapshots, page.url, network, notes
        finally:
            page.close()

    def close(self):
        if self.context:
            self.context.close()
        if self.browser:
            self.browser.close()
        if self.manager:
            self.manager.stop()
        self.context = self.browser = self.manager = None


def configure_document(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    # Some bundled Word templates attach a coloured bottom border to Title.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [("Title", 26), ("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.keep_with_next = True
        if name.startswith('Heading'):
            doc.styles[name].paragraph_format.space_before = Pt(12)
            doc.styles[name].paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def hyperlink(paragraph, label, url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rel)
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "245A81"); props.append(color)
    run.append(props); txt = OxmlElement("w:t"); txt.text = label; run.append(txt)
    link.append(run); paragraph._p.append(link)


def write_report(bundle, path, max_words=0):
    doc = Document(); configure_document(doc)
    doc.add_heading("Aqua Cloud Product Handbook", 0)
    doc.add_paragraph("Product capabilities and implementation reference", "Subtitle")
    doc.add_paragraph("This handbook groups collected public product and documentation content into a practical reading order. Begin with the overview, then follow the areas relevant to your QA workflow. Each source entry links back to its original page.")
    doc.add_paragraph("General chapter introductions explain testing concepts. The source detail is extracted website text, not an independent product assessment. Pricing, security, performance and planned-feature statements retain the status of vendor claims. No language model is used to fill gaps or invent explanations.")
    successful = bundle["pages"]
    doc.add_paragraph(f"Collected at {bundle['collected_at']} | Pages with usable content {len(successful)} | Attempted URLs {bundle['attempted']} | Pending URLs {bundle['pending_count']}")
    doc.add_heading("Reading order", 1)
    for i, (name, _, _) in enumerate(TOPICS, 1):
        count = sum(p["topic"] == i - 1 for p in successful)
        doc.add_paragraph(f"{i}. {name} — {count} source pages")
    doc.add_heading("Coverage and limitations", 1)
    doc.add_paragraph("This is a bounded crawl of accessible public pages, not proof of complete product coverage. Authenticated screens, image-only text, video narration, downloads and blocked or unvisited pages are not included. HTML may contain inactive tabs or hidden text; extraction is not a confirmation that a feature is visible or currently released. Third-party embedded tours require separate inspection.")
    if max_words:
        doc.add_paragraph(f"Source entries in this document are limited to {max_words} words each. The JSON retains the collected blocks. Truncated entries are marked.")
    for i, (name, _, intro) in enumerate(TOPICS):
        doc.add_heading(f"{i + 1} {name}", 1)
        doc.add_paragraph(intro)
        group = [p for p in successful if p["topic"] == i]
        if not group:
            doc.add_paragraph("No source page was assigned to this chapter in this run. Related material may appear in a broader overview page. This is a coverage gap, not evidence that aqua lacks the capability.")
        for page in group:
            doc.add_heading(page["title"], 2)
            p = doc.add_paragraph(f"[{page['id']}] ")
            hyperlink(p, page["url"], page["url"])
            doc.add_paragraph(f"Method {page['method']} | Retrieved {page['retrieved_at']} | {page['characters']} extracted characters")
            for note in page.get("notes", []):
                doc.add_paragraph("Collection note: " + note)
            remaining = max_words if max_words else float("inf")
            for block in page["blocks"]:
                words = block["text"].split()
                if remaining <= 0:
                    doc.add_paragraph("Source detail truncated by the configured word limit. See the collected JSON for all blocks.")
                    break
                text = " ".join(words[:int(remaining)]) if remaining != float("inf") else block["text"]
                remaining -= len(words)
                # Don't leave dangling headings at a source word limit.
                if block["kind"] == "heading" and remaining > 0:
                    doc.add_heading(text, 3)
                else:
                    doc.add_paragraph(text)
            if page.get("embeds"):
                doc.add_paragraph("Embedded media or tours were detected. Their internal content has not been transcribed.")
    doc.add_heading("Collection issues", 1)
    if not bundle["issues"]:
        doc.add_paragraph("No collection errors were recorded for attempted URLs. The page and depth limits still apply.")
    for item in bundle["issues"]:
        doc.add_paragraph(f"{item['url']} — {item['reason']}")
    doc.add_heading("Source register", 1)
    for page in successful:
        p = doc.add_paragraph(f"[{page['id']}] {page['title']} — ")
        hyperlink(p, page["url"], page["url"])
    doc.save(path)


def save_json(path, obj):
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def crawl(args):
    output = Path(args.output).resolve(); output.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", handlers=[logging.StreamHandler(), logging.FileHandler(output / "scrape.log", encoding="utf-8")], force=True)
    fetcher = Fetcher(args.delay, args.timeout, args.retries)
    browser = BrowserRenderer(args, fetcher)
    seeds = args.seed or SEEDS
    queue, queued, visited, hashes = deque(), set(), set(), {}
    for seed in seeds:
        u = normalise(seed)
        if not u:
            raise ValueError(f"Seed outside allowed public scope: {seed}")
        if u not in queued:
            queue.append((u, 0)); queued.add(u)
    bundle = {"schema_version": 1, "tool_version": VERSION, "collected_at": datetime.now(timezone.utc).isoformat(),
              "scope": {"hosts": sorted(HOSTS), "seeds": seeds, "mode": args.mode, "max_pages": args.max_pages, "max_depth": args.max_depth},
              "pages": [], "issues": [], "indexes": [], "network": [], "attempted": 0, "pending_count": 0, "robots": {}}
    try:
        while queue and bundle["attempted"] < args.max_pages:
            url, depth = queue.popleft()
            if url in visited:
                continue
            visited.add(url); bundle["attempted"] += 1
            logging.info("[%s/%s] %s", bundle["attempted"], args.max_pages, url)
            try:
                if not fetcher.allowed(url):
                    raise FetchError("Robots disallow or robots unavailable; see robots status")
                notes, network, parsed, method = [], [], None, "http"
                if args.mode != "browser":
                    raw, final_url, status, content_type = fetcher.raw(url)
                    parsed, method = extract_response(raw, final_url, content_type)
                    if final_url.endswith('/llms.txt'):
                        if depth < args.max_depth:
                            for link in parsed['links']:
                                if link not in queued and link not in visited and relevant(link):
                                    queue.append((link, depth + 1)); queued.add(link)
                        bundle['indexes'].append({'url': final_url, 'discovered_links': len(parsed['links'])})
                        continue
                # GitBook has valid, short navigation pages. Keep their links and
                # prose instead of mistaking an explicit Markdown response for a SPA.
                source_min = min(args.min_chars, 120) if method == 'http Markdown' else args.min_chars
                need_browser = args.mode == "browser" or (args.mode == "auto" and parsed["characters"] < source_min)
                if need_browser:
                    try:
                        snapshots, final_url, network, notes = browser.render(url)
                        rendered = extract_html(snapshots[0], final_url)
                        seen = {b["text"] for b in rendered["blocks"]}
                        for snapshot in snapshots[1:]:
                            state = extract_html(snapshot, final_url)
                            for block in state["blocks"]:
                                if block["text"] not in seen:
                                    rendered["blocks"].append(block); seen.add(block["text"])
                        txt = "\n".join(b["text"] for b in rendered["blocks"])
                        rendered["characters"] = len(txt)
                        rendered["content_hash"] = hashlib.sha256(txt.encode()).hexdigest()
                        if parsed is None or rendered["characters"] >= parsed["characters"]:
                            parsed, method = rendered, "browser"
                    except Exception as e:
                        if parsed is None:
                            raise FetchError(str(e)) from e
                        notes.append(str(e))
                bundle["network"].extend({"page": url, **n} for n in network)
                if parsed["characters"] < source_min:
                    raise FetchError(f"Insufficient content: {parsed['characters']} characters. " + "; ".join(notes))
                challenge = " ".join(b["text"] for b in parsed["blocks"][:4]).lower()
                if any(marker in challenge for marker in ["verify you are human", "checking your browser", "access denied", "enable javascript and cookies to continue"]):
                    raise FetchError("Possible access challenge; no bypass attempted")
                # Discover before deduplication so duplicate wrappers cannot cut off useful links.
                if depth < args.max_depth:
                    for link in sorted(parsed["links"], key=lambda u: (topic_for(u), len(u), u)):
                        if link not in queued and link not in visited and relevant(link):
                            queue.append((link, depth + 1)); queued.add(link)
                if parsed["content_hash"] in hashes:
                    bundle["issues"].append({"url": url, "reason": "Duplicate content of " + hashes[parsed["content_hash"]]})
                    continue
                hashes[parsed["content_hash"]] = parsed["url"]
                parsed.update({"id": f"S{len(bundle['pages']) + 1:03d}", "requested_url": url, "depth": depth, "method": method,
                               "topic": topic_for(url), "notes": notes, "retrieved_at": datetime.now(timezone.utc).isoformat()})
                bundle["pages"].append(parsed)
            except (FetchError, etree.ParserError, ValueError) as e:
                logging.warning("%s: %s", url, e)
                bundle["issues"].append({"url": url, "reason": str(e)})
            finally:
                bundle["pending_count"] = len(queue)
                bundle["robots"] = fetcher.robot_status
                save_json(output / "aqua_data.json", bundle)
    except KeyboardInterrupt:
        bundle["issues"].append({"url": "run", "reason": "Interrupted by user; output is partial"})
    finally:
        browser.close()
        bundle["pending_count"] = len(queue)
        bundle["pending_urls"] = [url for url, _ in queue]
        bundle["robots"] = fetcher.robot_status
        save_json(output / "aqua_data.json", bundle)
    if not bundle["pages"]:
        logging.error("No usable pages. See aqua_data.json and scrape.log. No handbook generated.")
        return 2
    write_report(bundle, output / "Aqua_Cloud_Handbook.docx", args.max_words_per_page)
    logging.info("Created %s from %s source pages; %s issues; %s queued URLs remain", output / "Aqua_Cloud_Handbook.docx", len(bundle["pages"]), len(bundle["issues"]), len(queue))
    return 0


def arguments(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--output", default="output", help="Output directory; files are replaced on rerun")
    p.add_argument("--mode", choices=["auto", "http", "browser"], default="auto")
    p.add_argument("--seed", action="append", help="Allowed public seed URL; repeat for multiple seeds")
    p.add_argument("--max-pages", type=int, default=50, help="Maximum attempted URLs, including failures")
    p.add_argument("--max-depth", type=int, default=2)
    p.add_argument("--delay", type=float, default=1.0, help="Minimum seconds between requests per host")
    p.add_argument("--timeout", type=int, default=30)
    p.add_argument("--retries", type=int, default=2)
    p.add_argument("--min-chars", type=int, default=400, help="Minimum extracted characters for usable content")
    p.add_argument("--max-words-per-page", type=int, default=0, help="Word-document limit per source; 0 means unlimited")
    p.add_argument("--headed", action="store_true", help="Show browser in browser mode")
    p.add_argument("--consent-selector", help="Inspected CSS selector for your chosen consent action; optional")
    p.add_argument("--next-selector", help="Inspected CSS selector for a tour next button; optional")
    p.add_argument("--max-steps", type=int, default=20)
    p.add_argument("--from-json", help="Rebuild the Word document from a previous aqua_data.json; no network")
    args = p.parse_args(argv)
    if args.max_pages < 1 or args.max_depth < 0 or args.delay < 0 or args.timeout < 1 or args.retries < 0 or args.min_chars < 1 or args.max_words_per_page < 0 or args.max_steps < 1:
        p.error("Invalid limits: page, time, content and step limits must be positive; depth, delay, retries and word cap cannot be negative")
    return args


def main(argv=None):
    args = arguments(argv)
    if args.from_json:
        bundle = json.loads(Path(args.from_json).read_text(encoding="utf-8"))
        output = Path(args.output); output.mkdir(parents=True, exist_ok=True)
        write_report(bundle, output / "Aqua_Cloud_Handbook.docx", args.max_words_per_page)
        print(output / "Aqua_Cloud_Handbook.docx")
        return 0
    return crawl(args)


if __name__ == "__main__":
    sys.exit(main())
